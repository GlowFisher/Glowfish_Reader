# -*- coding: utf-8 -*-
import io

from docx import Document
from docx.shared import Pt
from PIL import Image

import proof_model as M


def _annotated_proof():
    return {
        "batch": "sample",
        "script": "s",
        "notes": "research note",
        "pages": [{
            "page_id": "p1",
            "page_annot": "page note",
            "annots": [
                {"id": "a1", "quote": "battle", "occ": 0, "note": "first selection note"},
                {"id": "a2", "quote": "battle", "occ": 1, "note": "second selection note"},
                {"id": "a3", "quote": "missing", "occ": 0, "note": "missing selection note"},
            ],
            "lines": [
                {"id": "l1", "text": "year battle starts ", "deleted": False, "para_break": False, "annot": "line note"},
                {"id": "l2", "text": "battle spreads", "deleted": False, "para_break": True},
            ],
            "order": ["l1", "l2"],
        }],
    }


def test_docx_with_annotations_is_single_document_with_superscript_refs():
    data, count = M.doc_bytes(_annotated_proof(), fmt="docx", with_annotations=True)

    assert count == 1
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "year battle starts" in text
    assert "[1]" in text
    assert "[2]" in text
    assert "[3]" in text
    assert "first selection note" in text
    assert "line note" in text
    assert "second selection note" in text
    assert "missing selection note" in text
    assert "page note" in text
    assert "research note" in text

    ref_runs = [r for p in doc.paragraphs for r in p.runs if r.text == "[1]"]
    assert ref_runs and ref_runs[0].font.superscript


def test_docx_annotation_note_paragraphs_have_no_space_after():
    data, _count = M.doc_bytes(_annotated_proof(), fmt="docx", with_annotations=True)

    doc = Document(io.BytesIO(data))
    note_paras = [p for p in doc.paragraphs if p.text.startswith("[")]

    assert note_paras
    assert all(p.paragraph_format.space_after == Pt(0) for p in note_paras)


def test_markdown_with_annotations_uses_one_body_with_notes_section():
    data, count = M.doc_bytes(_annotated_proof(), fmt="md", with_annotations=True)

    body = data.decode("utf-8")
    assert count == 1
    assert "year battle[1] starts" in body
    assert "first selection note" in body
    assert "line note" in body
    assert "second selection note" in body
    assert "missing selection note" in body
    assert "page note" in body
    assert "research note" in body


def _image_proof(tmp_path):
    img1 = tmp_path / "p1.png"
    img2 = tmp_path / "p2.png"
    Image.new("RGB", (40, 20), "white").save(img1)
    Image.new("RGB", (30, 30), "black").save(img2)
    return {
        "batch": "sample",
        "script": "s",
        "notes": "document note",
        "pages": [
            {
                "page_id": "p1",
                "image": str(img1),
                "page_annot": "page one note",
                "annots": [{"quote": "page one", "occ": 0, "note": "selection note"}],
                "lines": [
                    {"id": "p1_l1", "text": "page one text", "deleted": False, "para_break": True, "annot": "line note"}
                ],
                "order": ["p1_l1"],
            },
            {
                "page_id": "p2",
                "image": str(img2),
                "lines": [
                    {"id": "p2_l1", "text": "page two text", "deleted": False, "para_break": True}
                ],
                "order": ["p2_l1"],
            },
        ],
    }


def test_docx_with_images_embeds_each_page_image_before_page_text(tmp_path):
    data, count = M.doc_bytes(_image_proof(tmp_path), fmt="docx", with_images=True)

    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)

    assert count == 2
    assert len(doc.inline_shapes) == 2
    assert "page one text" in text
    assert "page two text" in text
    assert text.index("page one text") < text.index("page two text")


def test_docx_with_images_and_annotations_keeps_notes_in_same_document(tmp_path):
    data, count = M.doc_bytes(
        _image_proof(tmp_path),
        fmt="docx",
        with_images=True,
        with_annotations=True,
    )

    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)

    assert count == 2
    assert len(doc.inline_shapes) == 2
    assert "page one text" in text
    assert "selection note" in text
    assert "line note" in text
    assert "page one note" in text
    assert "document note" in text

def test_docx_with_images_resolves_relative_image_against_image_dir(tmp_path):
    img_dir = tmp_path / "images"
    img_dir.mkdir()
    Image.new("RGB", (24, 24), "white").save(img_dir / "p1.png")
    proof = {
        "batch": "sample",
        "script": "s",
        "image_dir": str(img_dir),
        "pages": [{
            "page_id": "p1",
            "image": "p1.png",
            "lines": [{"id": "p1_l1", "text": "relative image text", "deleted": False, "para_break": True}],
            "order": ["p1_l1"],
        }],
    }

    data, count = M.doc_bytes(proof, fmt="docx", with_images=True)

    doc = Document(io.BytesIO(data))
    assert count == 1
    assert len(doc.inline_shapes) == 1
    assert "relative image text" in "\n".join(p.text for p in doc.paragraphs)
