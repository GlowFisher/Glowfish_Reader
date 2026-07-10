# -*- coding: utf-8 -*-
"""校对数据模型与拼装/导出（codex env，无 paddle 依赖）。

三稿纪律：text_draft(初稿只读) → text(人工修订) → 导出定稿。
核心可测函数：build_proof / assemble_paragraphs / export_docx / edit_stats。
"""
import sys
import os
import re
import json
import unicodedata
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "pipeline"))
from build_docx import make_doc  # noqa: E402
from opencc import OpenCC  # noqa: E402

import t2s_fixes  # noqa: E402  OpenCC t2s 漏网字修正层（日式简化字/舖等，纯数据表）
import interlinear  # noqa: E402  迭代1 双行夹注读序（纯几何，挂 aggregate_columns 之前）
import domain_fixes  # noqa: E402  A1 域词确定性修复层（纯数据表，空跑审计准入）
import ocr_correct  # noqa: E402  迭代3 上下文语义纠错（顶层无 onnxruntime 依赖，惰性加载）
import name_segment  # noqa: E402  迭代3e 并列人名句读建议（Tier B，纯几何/姓氏，无重依赖）

LOW_CONF = 0.85  # 低置信高亮阈值
_T2S = OpenCC("t2s")  # 繁→简
_S2T = OpenCC("s2t")  # 简→繁

# ---- 2b：纪年非法数字字纠正（确定性规则，纪年语境触发，保守）----
_YEAR_NUMERAL = set("一二三四五六七八九十百千零〇廿卅两")
# 只把「本身非合法数字」的形近字补成数字（如 川→十）；本身合法的数字（七等）不动——
# 七既可能是 7、也可能是误认的 1，属歧义，交迭代3 MLM 上下文裁决（「置换后序列合法」在歧义下的保守落地）。
_YEAR_INVALID_CONFUSION = {"川": "十", "什": "十", "卄": "廿", "仟": "千", "佰": "百"}
_REIGN_RE = re.compile(
    r"(中华民国|大清|民国|光绪|宣统|康熙|乾隆|雍正|洪武|崇祯|同治|咸丰|道光|嘉庆)"
    r"([一二三四五六七八九十百千零〇廿卅两川什卄仟佰]{1,6})年"
)


def _fix_reign_year(text):
    """朝代语境下的纪年数字槽：非法数字字→形近合法数字，且要求替换后全为合法数字。"""
    if not text or "年" not in text:
        return text

    def repl(m):
        prefix, slot = m.group(1), m.group(2)
        fixed = "".join(_YEAR_INVALID_CONFUSION.get(c, c) for c in slot)
        if fixed != slot and all(c in _YEAR_NUMERAL for c in fixed):
            return prefix + fixed + "年"
        return m.group(0)

    return _REIGN_RE.sub(repl, text)


def _reign_year_enabled():
    return os.environ.get("SHUPAI_REIGN_YEAR", "1").strip().lower() not in ("0", "false", "no", "")


# ---- 关键词搜索：简繁互通 + 忽略标点（点校区搜索用）----
def _norm_for_search(s):
    """繁→简 + 去标点/空白/控制符；返回 (norm, idx_map)，idx_map[i]=norm 第 i 字在原串 s 的下标。

    逐字转换以保持「norm 字符 → 原文下标」可逆映射，便于把命中位置映射回原文坐标。"""
    norm, idx_map = [], []
    for i, ch in enumerate(s or ""):
        for cc in _T2S.convert(ch):                 # 单字繁→简（极少 1:多，整体仍映射到同一原文 i）
            if unicodedata.category(cc)[0] in ("P", "Z", "C"):   # 标点 / 分隔(空白) / 控制 → 丢弃
                continue
            norm.append(cc)
            idx_map.append(i)
    return "".join(norm), idx_map


def search_in_texts(query, texts):
    """在每条 text 里找 query 的所有出现（简繁互通 + 忽略标点）。

    返回与 texts 等长的列表，每项为该 text 的命中区间 [[start, end], ...]（原文坐标，end 不含）。
    空 query 返回各空列表；命中不重叠。"""
    qn, _ = _norm_for_search(query or "")
    out = []
    for t in (texts or []):
        ranges = []
        if qn:
            tn, idx_map = _norm_for_search(t)
            start = 0
            while True:
                j = tn.find(qn, start)
                if j < 0:
                    break
                s0 = idx_map[j]
                e0 = idx_map[j + len(qn) - 1] + 1   # 末命中字原文下标 + 1
                ranges.append([s0, e0])
                start = j + len(qn)                 # 不重叠
        out.append(ranges)
    return out


def _mean_x(box):
    return (box[0] + box[2]) / 2.0


def _mean_y(box):
    return (box[1] + box[3]) / 2.0


def default_order(lines):
    """右起栏序：列(x)降序、列内(y)升序。返回 line id 列表。"""
    idx = sorted(range(len(lines)),
                 key=lambda i: (-_mean_x(lines[i]["box"]), _mean_y(lines[i]["box"])))
    return [lines[i]["id"] for i in idx]


def _x_overlap_ratio(a, b):
    """两 box 在 x 方向的重叠占较窄宽度的比例（竖列聚类用）。"""
    inter = max(0.0, min(a[2], b[2]) - max(a[0], b[0]))
    narrow = max(1.0, min(a[2] - a[0], b[2] - b[0]))
    return inter / narrow


def _y_overlap_ratio(a, b):
    """两 box 在 y 方向的重叠占较矮高度的比例。"""
    inter = max(0.0, min(a[3], b[3]) - max(a[1], b[1]))
    short = max(1.0, min(a[3] - a[1], b[3] - b[1]))
    return inter / short


def _col_agg_yguard_enabled():
    return os.environ.get("SHUPAI_COL_AGG_YGUARD", "1").strip().lower() not in ("0", "false", "no", "")


def _same_column(a, b):
    """能否判定 a、b 属于同一竖列。

    同列碎片是**上下堆叠**的（y 基本不重叠）；行草页里相邻两列的框常因笔画倾斜
    在 x 上重叠 ≥0.5，只看 x 会被并查集链式并成一列，两列文字按 y 交错拼接、
    读序全乱（实测退想斋日记页）。故再加一道 y 守卫。
    """
    if _x_overlap_ratio(a, b) < 0.5:
        return False
    if _col_agg_yguard_enabled() and _y_overlap_ratio(a, b) >= 0.30:
        return False
    return True


def aggregate_columns(raw_lines):
    """竖行兜底聚合：把同一竖列被 OCR 切碎的多块合并为单块。"""
    n = len(raw_lines)
    boxes = [[float(x) for x in l.get("box", [0, 0, 0, 0])] for l in raw_lines]
    if n <= 1:
        out = []
        for i, l in enumerate(raw_lines):
            item = {"text": l.get("text", ""), "box": boxes[i], "score": float(l.get("score", 0.0))}
            for meta_key in ("source", "filter_status", "filter_reason"):
                if meta_key in l:
                    item[meta_key] = l[meta_key]
            out.append(item)
        return out
    parent = list(range(n))

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]; i = parent[i]
        return i

    for i in range(n):
        for j in range(i + 1, n):
            if _same_column(boxes[i], boxes[j]):
                ri, rj = find(i), find(j)
                if ri != rj:
                    parent[ri] = rj
    groups = {}
    for i in range(n):
        groups.setdefault(find(i), []).append(i)
    merged = []
    for idxs in sorted(groups.values(), key=lambda g: min(g)):
        idxs.sort(key=lambda k: (boxes[k][1] + boxes[k][3]) / 2.0)
        if len(idxs) == 1:
            k = idxs[0]
            item = {"text": raw_lines[k].get("text", ""), "box": boxes[k],
                    "score": float(raw_lines[k].get("score", 0.0))}
            for meta_key in ("source", "filter_status", "filter_reason"):
                if meta_key in raw_lines[k]:
                    item[meta_key] = raw_lines[k][meta_key]
            merged.append(item)
            continue
        text = "".join(raw_lines[k].get("text", "") for k in idxs)
        bx = [min(boxes[k][0] for k in idxs), min(boxes[k][1] for k in idxs),
              max(boxes[k][2] for k in idxs), max(boxes[k][3] for k in idxs)]
        item = {"text": text, "box": bx,
                "score": min(float(raw_lines[k].get("score", 0.0)) for k in idxs),
                "member_boxes": [boxes[k] for k in idxs],
                "member_texts": [raw_lines[k].get("text", "") for k in idxs]}
        sources = sorted({str(raw_lines[k].get("source", "")) for k in idxs if raw_lines[k].get("source")})
        if sources:
            item["source"] = "+".join(sources)
        merged.append(item)
    return merged


def _proof_line(pid, j, m):
    box = m["box"]
    t = m["text"]
    ts = t2s_fixes.apply(_T2S.convert(t))
    if _reign_year_enabled():
        ts = _fix_reign_year(ts)  # 2b：纪年非法数字字纠正（纪年语境，保守）
    ln = {
        "id": f"{pid}__{j}",
        "text_ocr": t,
        "text_draft": ts,
        "text": ts,
        "box": box,
        "score": float(m.get("score", 0.0)),
        "para_break": False,
        "deleted": False,
    }
    for meta_key in ("source", "filter_status", "filter_reason"):
        if meta_key in m:
            ln[meta_key] = m[meta_key]
    if "member_texts" in m:
        ln["member_boxes"] = m["member_boxes"]
        ln["member_texts"] = m["member_texts"]
    return ln


def _noise_line(pid, j, m):
    ln = _proof_line(pid + "__noise", j, m)
    ln["id"] = f"{pid}__noise__{j}"
    ln["deleted"] = True
    return ln


def build_proof(recog, batch_name="batch", engine="paddle"):
    """富识别 recog dict -> proof dict.

    OCR output is split into editable body lines and a recoverable noise layer.
    Noise lines stay in proof.json but are not part of the default order/export.
    """
    pages = []
    for pid, pg in recog.items():
        if str(pid).startswith("_"):
            continue
        raw = pg.get("lines", [])
        main_raw = [l for l in raw if l.get("filter_status") != "noise"]
        main_raw = interlinear.pair_and_merge(main_raw)  # 迭代1：双行夹注对合并为合成行（宁漏勿错）
        noise_raw = [l for l in raw if l.get("filter_status") == "noise"]
        lines = [_proof_line(pid, j, m) for j, m in enumerate(aggregate_columns(main_raw))]
        try:
            lines = domain_fixes.correct_page(lines)  # A1：域词确定性修复（空跑审计证明零误纠），先于 MLM
        except Exception:
            pass
        try:
            lines = ocr_correct.correct_page(lines)  # 迭代3：MLM 语义纠错（任何异常静默降级）
        except Exception:
            pass
        try:
            lines = name_segment.annotate_lines(lines)  # 迭代3e：并列人名句读建议（Tier B，非名单语境零触发）
        except Exception:
            pass
        noise_lines = [_noise_line(pid, j, m) for j, m in enumerate(noise_raw)]
        pages.append({
            "page_id": pid,
            "image": pg.get("image", f"{pid}.png"),
            "width": pg.get("width", 0),
            "height": pg.get("height", 0),
            "lines": lines,
            "noise_lines": noise_lines,
            "ocr_meta": pg.get("ocr_meta", {}),
            "order": default_order(lines),
        })
    return {"batch": batch_name, "engine": engine, "script": "s", "pages": pages}


def convert_proof(proof, to):
    """简繁在位转换所有 text。to='s'(简) 或 't'(繁)。设 proof['script']。
    简向输出过 t2s_fixes（修 OpenCC 漏网的日式简化字等）；繁向不动。"""
    cc = _T2S if to == "s" else _S2T
    for pg in proof["pages"]:
        for l in pg["lines"]:
            t = cc.convert(l.get("text", ""))
            l["text"] = t2s_fixes.apply(t) if to == "s" else t
    proof["script"] = to
    return proof


def _page_label(page_id, ordinal):
    """页码标注文本。优先取 page_id 尾部 _pN（PDF 源页号），无则用 proof 内 1-based 顺序号。"""
    m = re.search(r"_p(\d+)$", page_id or "")
    n = int(m.group(1)) if m else ordinal
    return f"【第{n}页】"


def assemble_paragraphs(proof, page_marks=False):
    """按 order + para_break + deleted 拼成段落列表（用 text 修订稿）。

    page_marks=True 时，在每页正文前插入独立的「【第N页】」标注段（跨页段落在页界处断开）。
    默认 False 时行为与原先完全一致（段落可跨页连续）。"""
    paras, cur = [], []
    for ordinal, pg in enumerate(proof["pages"], 1):
        lmap = {l["id"]: l for l in pg["lines"]}
        order_ids = [lid for lid in pg.get("order", [])
                     if lmap.get(lid) and not lmap[lid].get("deleted")]
        if page_marks and order_ids:
            if cur:
                paras.append("".join(cur))
                cur = []
            paras.append(_page_label(pg.get("page_id", ""), ordinal))
        for lid in order_ids:
            cur.append(lmap[lid].get("text", ""))
            if lmap[lid].get("para_break"):
                paras.append("".join(cur))
                cur = []
    if cur:
        paras.append("".join(cur))
    return [p for p in paras if p.strip()]


def export_docx(proof, out_path, t2s=True, blank_between=True, page_marks=False):
    """导出 docx。blank_between=True 时自然段之间自动空一行（插入空段）。
    page_marks=True 时每页正文前插「【第N页】」。"""
    paras = assemble_paragraphs(proof, page_marks=page_marks)
    n = len(paras)
    if blank_between and n > 1:
        spaced = []
        for i, p in enumerate(paras):
            if i:
                spaced.append("")  # 自然段间空行
            spaced.append(p)
        paras = spaced
    make_doc(paras, out_path, t2s=t2s)
    return out_path, n


# 导出格式 → 文件后缀
EXT = {"docx": ".docx", "txt": ".txt", "md": ".md", "csv": ".csv"}


def _paras_to_bytes(paras, name, fmt):
    """非 docx 段落 → 编码字节（txt/md/csv）。csv 用 utf-8-sig 让 Excel 不乱码。"""
    if fmt == "txt":
        body = "\n\n".join(paras)
    elif fmt == "md":
        body = ("# " + name + "\n\n") + "\n\n".join(paras)
    else:  # csv
        import csv as _csv
        import io as _io
        buf = _io.StringIO()
        w = _csv.writer(buf)
        w.writerow(["序", "内容"])
        for i, p in enumerate(paras, 1):
            w.writerow([i, p])
        body = buf.getvalue()
    return body.encode("utf-8-sig" if fmt == "csv" else "utf-8")



def _nth_find(text, needle, occ):
    if not needle:
        return -1
    start = 0
    for _ in range(max(0, int(occ or 0)) + 1):
        pos = text.find(needle, start)
        if pos < 0:
            return -1
        start = pos + len(needle)
    return pos


def _page_annotated_segments(pg):
    """Return page body segments and footnote-like notes."""
    lmap = {l["id"]: l for l in pg.get("lines", [])}
    lines = [lmap[lid] for lid in pg.get("order", [])
             if lmap.get(lid) and not lmap[lid].get("deleted")]
    if not lines:
        return [], []
    text = "".join(l.get("text", "") for l in lines)
    notes = []

    def add_marker(pos, label, note):
        note = (note or "").strip()
        if note:
            notes.append({"pos": max(0, min(len(text), pos)), "label": label, "note": note})

    for a in pg.get("annots", []) or []:
        q = (a.get("quote") or "").strip()
        pos = _nth_find(text, q, a.get("occ", 0))
        add_marker((pos + len(q)) if pos >= 0 else len(text), q or "失锚批注", a.get("note"))

    offset = 0
    for l in lines:
        lt = l.get("text", "")
        add_marker(offset + len(lt), lt[:24] or "段落", l.get("annot"))
        offset += len(lt)
    add_marker(len(text), "本页", pg.get("page_annot"))

    if not notes:
        return [[(text, False)]], []
    notes.sort(key=lambda n: n["pos"])
    segs, cursor = [], 0
    for idx, n in enumerate(notes, 1):
        pos = n["pos"]
        if pos > cursor:
            segs.append((text[cursor:pos], False))
        segs.append((f"[{idx}]", True))
        n["num"] = idx
        cursor = pos
    if cursor < len(text):
        segs.append((text[cursor:], False))
    return [segs], notes


def annotated_paragraphs(proof, page_marks=False):
    body, notes, para_count, next_num = [], [], 0, 1
    for ordinal, pg in enumerate(proof.get("pages", []), 1):
        page_body, page_notes = _page_annotated_segments(pg)
        if not page_body:
            continue
        if page_marks:
            body.append([(_page_label(pg.get("page_id", ""), ordinal), False)])
        for segs in page_body:
            shifted = []
            for text, sup in segs:
                if sup:
                    local_num = int(str(text).strip("[]"))
                    shifted.append((f"[{next_num + local_num - 1}]", True))
                else:
                    shifted.append((text, False))
            body.append(shifted)
            para_count += 1
        for n in page_notes:
            notes.append(f"[{next_num}] {n['label']}：{n['note']}")
            next_num += 1
    doc_notes = (proof.get("notes") or "").strip()
    if doc_notes:
        notes.append(f"[{next_num}] 研究笔记：{doc_notes}")
    return body, notes, para_count


def _annotated_docx_bytes(proof, t2s=False, blank_between=True, page_marks=False):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.oxml.ns import qn
    import io as _io

    body, notes, n = annotated_paragraphs(proof, page_marks=page_marks)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(3.2)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    font, size_pt = "宋体", 10.5
    normal = doc.styles["Normal"]
    normal.font.size = Pt(size_pt)
    normal.font.name = font
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    indent = Pt(size_pt * 2)

    def add_para(segments, first_indent=True, small=False):
        p = doc.add_paragraph()
        if first_indent:
            p.paragraph_format.first_line_indent = indent
        if small:
            p.paragraph_format.space_after = Pt(0)
        for text, sup in segments:
            run = p.add_run(_T2S.convert(text) if t2s else text)
            run.font.size = Pt(9 if small else size_pt)
            run.font.name = font
            run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
            if sup:
                run.font.superscript = True

    for i, segs in enumerate(body):
        if blank_between and i:
            doc.add_paragraph()
        add_para(segs)
    if notes:
        doc.add_paragraph()
        head = doc.add_paragraph()
        r = head.add_run("注释")
        r.bold = True
        r.font.name = font
        r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        for note in notes:
            add_para([(note, False)], first_indent=False, small=True)
    buf = _io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), n



def _docx_base():
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.oxml.ns import qn

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(3.2)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    font, size_pt = "\u5b8b\u4f53", 10.5
    normal = doc.styles["Normal"]
    normal.font.size = Pt(size_pt)
    normal.font.name = font
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return doc, sec, font, size_pt, Pt(size_pt * 2), qn


def _add_docx_segments(doc, segments, font, size_pt, qn, indent=None, small=False):
    from docx.shared import Pt

    p = doc.add_paragraph()
    if indent is not None:
        p.paragraph_format.first_line_indent = indent
    if small:
        p.paragraph_format.space_after = Pt(0)
    for text, sup in segments:
        run = p.add_run(text)
        run.font.size = Pt(9 if small else size_pt)
        run.font.name = font
        run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        if sup:
            run.font.superscript = True
    return p



def _page_image_path(proof, pg):
    raw = str(pg.get("image", ""))
    src = Path(raw)
    image_dir = proof.get("image_dir")
    candidates = []
    if src.is_absolute():
        candidates.append(src)
    elif image_dir:
        candidates.append(Path(image_dir) / raw)
    elif raw:
        candidates.append(src)
    if image_dir:
        base = Path(image_dir)
        names = []
        if raw:
            names.append(Path(raw).name)
        page_id = str(pg.get("page_id") or "")
        if page_id:
            names.append(Path(page_id).name)
        for name in names:
            candidates.append(base / name)
            stem = Path(name).stem
            for ext in (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"):
                candidates.append(base / f"{stem}{ext}")
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return candidates[0] if candidates else src
def _add_page_image(doc, proof, pg, max_width):
    src = _page_image_path(proof, pg)
    if not src.exists():
        return False
    try:
        p = doc.add_paragraph()
        p.alignment = 1
        p.add_run().add_picture(str(src), width=max_width)
        return True
    except Exception:
        return False


def _image_docx_bytes(proof, t2s=False, page_marks=False, with_annotations=False):
    import io as _io

    doc, sec, font, size_pt, indent, qn = _docx_base()
    max_width = sec.page_width - sec.left_margin - sec.right_margin
    page_count, next_num = 0, 1

    for ordinal, pg in enumerate(proof.get("pages", []), 1):
        page_body, page_notes = _page_annotated_segments(pg)
        if not page_body:
            continue
        if page_count:
            doc.add_paragraph()
        _add_page_image(doc, proof, pg, max_width)
        if page_marks:
            _add_docx_segments(doc, [(_page_label(pg.get("page_id", ""), ordinal), False)], font, size_pt, qn, indent)
        if with_annotations:
            for segs in page_body:
                shifted = []
                for text, sup in segs:
                    if sup:
                        local_num = int(str(text).strip("[]"))
                        shifted.append((f"[{next_num + local_num - 1}]", True))
                    else:
                        shifted.append((_T2S.convert(text) if t2s else text, False))
                _add_docx_segments(doc, shifted, font, size_pt, qn, indent)
            if page_notes:
                head = doc.add_paragraph()
                r = head.add_run("\u6ce8\u91ca")
                r.bold = True
                r.font.name = font
                r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
                for n in page_notes:
                    note = f"[{next_num}] {n['label']}\uff1a{n['note']}"
                    if t2s:
                        note = _T2S.convert(note)
                    _add_docx_segments(doc, [(note, False)], font, size_pt, qn, None, small=True)
                    next_num += 1
        else:
            paras = assemble_paragraphs({**proof, "pages": [pg]}, page_marks=False)
            for para in paras:
                text = _T2S.convert(para) if t2s else para
                _add_docx_segments(doc, [(text, False)], font, size_pt, qn, indent)
        page_count += 1

    doc_notes = (proof.get("notes") or "").strip()
    if with_annotations and doc_notes:
        doc.add_paragraph()
        head = doc.add_paragraph()
        r = head.add_run("\u6ce8\u91ca")
        r.bold = True
        r.font.name = font
        r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        note = f"[{next_num}] \u7814\u7a76\u7b14\u8bb0\uff1a{doc_notes}"
        if t2s:
            note = _T2S.convert(note)
        _add_docx_segments(doc, [(note, False)], font, size_pt, qn, None, small=True)

    buf = _io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), page_count

def _annotated_text_bytes(proof, name, fmt, page_marks=False):
    body, notes, n = annotated_paragraphs(proof, page_marks=page_marks)
    paras = ["".join(t for t, _sup in segs) for segs in body]
    if notes:
        paras.append("## 注释" if fmt == "md" else "注释")
        paras.extend(notes)
    return _paras_to_bytes(paras, name, fmt), n

def doc_bytes(proof, fmt="docx", t2s=False, blank_between=True, page_marks=False, with_annotations=False, with_images=False):
    """导出一份 proof 为内存字节（浏览器直接下载用）。返回 (data: bytes, 自然段数)。
    page_marks=True 时每页正文前插「【第N页】」。"""
    import io as _io
    fmt = fmt if fmt in EXT else "docx"
    if with_images and fmt == "docx":
        return _image_docx_bytes(proof, t2s=t2s, page_marks=page_marks, with_annotations=with_annotations)
    if with_annotations:
        if fmt == "docx":
            return _annotated_docx_bytes(proof, t2s=t2s, blank_between=blank_between, page_marks=page_marks)
        return _annotated_text_bytes(proof, proof.get("batch", "辑录"), fmt, page_marks=page_marks)
    paras = assemble_paragraphs(proof, page_marks=page_marks)
    n = len(paras)
    if fmt == "docx":
        out = paras
        if blank_between and n > 1:                  # 自然段间空一行
            out = []
            for i, p in enumerate(paras):
                if i:
                    out.append("")
                out.append(p)
        buf = _io.BytesIO()
        make_doc(out, buf, t2s=t2s)                  # python-docx save() 原生支持 BytesIO
        return buf.getvalue(), n
    return _paras_to_bytes(paras, proof.get("batch", "辑录"), fmt), n


def paras_doc_bytes(paras, name, fmt="docx"):
    """已就绪段落列表 → 内存字节（全本合并 / 仅批注导出复用）。返回 (data, 段数)。"""
    import io as _io
    fmt = fmt if fmt in EXT else "docx"
    if fmt == "docx":
        buf = _io.BytesIO()
        make_doc(paras, buf, t2s=False)
        return buf.getvalue(), len(paras)
    return _paras_to_bytes(paras, name, fmt), len(paras)


def _edit_distance(a, b):
    la, lb = len(a), len(b)
    if la == 0:
        return lb
    if lb == 0:
        return la
    prev = list(range(lb + 1))
    for i in range(1, la + 1):
        cur = [i] + [0] * lb
        ai = a[i - 1]
        for j in range(1, lb + 1):
            cur[j] = min(prev[j] + 1, cur[j - 1] + 1,
                         prev[j - 1] + (0 if ai == b[j - 1] else 1))
        prev = cur
    return prev[lb]


def edit_stats(proof):
    """省人工度量：初稿全文 vs 定稿全文的编辑距离与改动占比。"""
    draft = "".join(l["text_draft"] for pg in proof["pages"] for l in pg["lines"])
    final_raw = "".join(l["text"] for pg in proof["pages"] for l in pg["lines"]
                        if not l.get("deleted"))
    final = _T2S.convert(final_raw)  # 归一到简,只量人工编辑、不计简繁切换
    dist = _edit_distance(draft, final)
    base = max(1, len(final))
    return {
        "draft_chars": len(draft),
        "final_chars": len(final),
        "edit_distance": dist,
        "changed_ratio": round(dist / base, 4),
        "saved_ratio": round(1 - dist / base, 4),  # 大致“沿用了多少初稿”
    }


def load_or_build(proof_path, recog_path, batch_name, engine="paddle"):
    """有 proof.json 读之；否则从 recog 建并落盘。"""
    proof_path, recog_path = Path(proof_path), Path(recog_path)
    if proof_path.exists():
        proof = json.loads(proof_path.read_text(encoding="utf-8"))
        proof.setdefault("engine", engine)      # 旧档缺字段回填（向后兼容）
        return proof
    recog = json.loads(recog_path.read_text(encoding="utf-8"))
    proof = build_proof(recog, batch_name, engine)
    proof_path.parent.mkdir(parents=True, exist_ok=True)
    proof_path.write_text(json.dumps(proof, ensure_ascii=False, indent=2), encoding="utf-8")
    return proof


def save_proof(proof, proof_path):
    Path(proof_path).write_text(json.dumps(proof, ensure_ascii=False, indent=2), encoding="utf-8")
